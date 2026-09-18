import io
import re
from docx import Document

FIELD_RE = re.compile(r"{{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*}}")

def _paragraph_texts(document):
    chunks = []
    for p in document.paragraphs:
        chunks.append(p.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    chunks.append(p.text)

    for section in document.sections:
        for part in [section.header, section.footer]:
            for p in part.paragraphs:
                chunks.append(p.text)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            chunks.append(p.text)

    return chunks

def inspect_template(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    text = "\n".join(_paragraph_texts(doc))
    fields = sorted(set(FIELD_RE.findall(text)))
    return {"fields": fields, "text_length": len(text)}
